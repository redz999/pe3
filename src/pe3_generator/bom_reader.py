from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .models import BomRow

REQUIRED_COLUMNS = ["Designator", "Part Number", "Quantity"]


class BomValidationError(Exception):
    pass


def read_bom(path: str | Path) -> list[BomRow]:
    workbook = load_workbook(filename=path, data_only=True)
    sheet = workbook.worksheets[0]

    header_cells = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(x).strip() if x is not None else "" for x in header_cells]

    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        missing_text = ", ".join(missing)
        raise BomValidationError(
            f"Неправильный состав BOM-файла. Отсутствуют колонки: {missing_text}"
        )

    indices = {name: headers.index(name) for name in REQUIRED_COLUMNS}

    rows: list[BomRow] = []
    for raw_row in sheet.iter_rows(min_row=2, values_only=True):
        designator = raw_row[indices["Designator"]]
        part_number = raw_row[indices["Part Number"]]
        quantity = raw_row[indices["Quantity"]]

        if designator is None or part_number is None:
            continue

        qty_int = int(quantity) if quantity not in (None, "") else 1
        rows.append(
            BomRow(
                designator=str(designator).strip(),
                part_number=str(part_number).strip(),
                quantity=qty_int,
            )
        )

    return rows
